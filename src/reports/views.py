"""
Виды (views) для всех отчетов.
"""
import calendar
import json
import re
from datetime import datetime, timedelta
from io import BytesIO
from django.utils import timezone as dj_timezone

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Font, PatternFill
from django import forms
from django.contrib import messages
from django.http import HttpResponse, HttpResponseRedirect, Http404
from django.shortcuts import redirect, render
from django.urls import reverse, reverse_lazy
from django.utils.decorators import method_decorator
from django.views import View
from django.views.generic import TemplateView, CreateView, UpdateView

from accounts.views import admin_required
from production.models import LotInfo, ManualDefect, ProductionRecord, Report5Comment, SNComment
from reports.queries import (
    _aggregate_monthly,
    get_date_range,
    report1_orders,
    report2_line_productivity,
    report4_defects,
    report5_repeated_passes,
    report6_serial_number,
    report7_station130_output,
    report8_stations,
    report8_table_columns,
    REPORT8_COLUMNS,
    REPORT8_CHART_STATIONS,
    REPORT8_CHART_KEYS,
    REPORT8_CHART_COLORS,
    REPORT8_HEADER_GROUPS,
)


# ========================
# Сортировка
# ========================

SORTABLE_FIELDS = {
    "report1": ["LOT", "production_spec", "Total", "Production", "Ok",
                 "Diff", "Defect", "start", "finish", "comment"],
    "report2": ["report_date", "fol_hours", "fol_production", "fol_speed",
                 "bol_hours", "bol_production", "bol_speed"],
    "report3": ["report_month", "fol_hours", "fol_production", "fol_speed",
                 "bol_hours", "bol_production", "bol_speed"],
    "report4": ["pcs_no", "production_spec", "lot_number", "entry_date", "last_station", "last_station_date", "comment"],
    "report5": ["pcs_no", "production_spec", "lot_number", "station", "dates", "comment"],
    "report6": ["pcs_no", "lot_number", "production_spec", "subop_no", "workstation_name",
                "created_date", "result", "test_data", "comment"],
     "report7": ["pcs_no", "lot_number", "production_spec", "created_date"],
    "report8": ["report_date", "fol_diff", "bol_diff", "station_020_pct", "station_080_pct",
                "station_090_pct", "station_100_pct", "station_110_pct",
                "station_120_pct", "station_130_pct"],
}


def apply_sort(data, sort_field, sort_dir, allowed_fields):
    """
    Сортирует список словарей по полю. None всегда сортируются в конец.
    """
    if sort_field not in allowed_fields or not isinstance(data, list):
        return data

    def get_key(item):
        val = item.get(sort_field) if isinstance(item, dict) else None
        if val is None:
            return (1,)
        if isinstance(val, (int, float)):
            return (0, 0, val)
        if isinstance(val, str):
            return (0, 1, val.casefold())
        return (0, 2, val)

    none_items = [item for item in data if get_key(item)[0] == 1]
    non_none = [item for item in data if item not in none_items]
    non_none.sort(key=get_key, reverse=(sort_dir == "desc"))
    return non_none + none_items


def get_sort_params(request, report_key):
    """Извлекает и валидирует параметры сортировки из GET-запроса."""
    allowed = SORTABLE_FIELDS[report_key]
    sort_field = request.GET.get("sort", "")
    sort_dir = request.GET.get("dir", "asc")
    if sort_field not in allowed:
        sort_field = ""
    if sort_dir not in ("asc", "desc"):
        sort_dir = "asc"
    return sort_field, sort_dir


REPORT7_PERIOD_TYPES = {"day", "week", "month", "period"}


def _report7_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _report7_int(value, default, lower, upper):
    try:
        parsed = int(value)
        return parsed if lower <= parsed <= upper else default
    except (TypeError, ValueError):
        return default


def _report7_period(request):
    today = dj_timezone.localdate()
    filter_type = request.GET.get("filter", "day")
    if filter_type not in REPORT7_PERIOD_TYPES:
        filter_type = "day"

    iso_year, iso_week, _ = today.isocalendar()
    day_value = today.isoformat()
    week_value = f"{iso_year}-W{iso_week:02d}"
    month_value = today.month
    year_value = today.year
    start_date = today
    end_date = today
    error = ""

    if filter_type == "day":
        selected_day = _report7_date(request.GET.get("day")) or today
        day_value = selected_day.isoformat()
        iso_year, iso_week, _ = selected_day.isocalendar()
        week_value = f"{iso_year}-W{iso_week:02d}"
        month_value = selected_day.month
        year_value = selected_day.year
        start_date = selected_day
        end_date = selected_day
    elif filter_type == "week":
        week_value = request.GET.get("week") or week_value
        week_match = re.fullmatch(r"(\d{4})-W(\d{2})", week_value)
        if week_match:
            try:
                start_date = datetime.fromisocalendar(
                    int(week_match.group(1)), int(week_match.group(2)), 1
                ).date()
                end_date = start_date + timedelta(days=6)
                day_value = start_date.isoformat()
                month_value = start_date.month
                year_value = start_date.year
            except ValueError:
                error = "Выбрана некорректная неделя."
        else:
            error = "Выбрана некорректная неделя."
        if error:
            start_date = today
            end_date = today
    elif filter_type == "month":
        year_value = _report7_int(request.GET.get("year"), today.year, 2000, 2100)
        month_value = _report7_int(request.GET.get("month"), today.month, 1, 12)
        try:
            last_day = calendar.monthrange(year_value, month_value)[1]
            start_date = datetime(year_value, month_value, 1).date()
            end_date = datetime(year_value, month_value, last_day).date()
            day_value = start_date.isoformat()
            iso_year, iso_week, _ = start_date.isocalendar()
            week_value = f"{iso_year}-W{iso_week:02d}"
        except ValueError:
            error = "Выбран некорректный месяц."
            start_date = today
            end_date = today
    else:
        start_value = request.GET.get("start")
        end_value = request.GET.get("end")
        start_date = _report7_date(start_value)
        end_date = _report7_date(end_value)
        if start_date is None or end_date is None:
            error = "Выбран некорректный период."
            start_date = today
            end_date = today
        elif end_date < start_date:
            error = "Дата начала не может быть позже даты окончания."
            start_date = today
            end_date = today
        day_value = start_date.isoformat()

    return {
        "filter_type": filter_type,
        "start_date": start_date,
        "end_date": end_date,
        "day_value": day_value,
        "week_value": week_value,
        "month_value": month_value,
        "year_value": year_value,
        "error": error,
    }


REPORT2_PERIOD_TYPES = {"month", "year", "period", "lot"}


def _report2_int(value, default, lower, upper):
    try:
        parsed = int(value)
        return parsed if lower <= parsed <= upper else default
    except (TypeError, ValueError):
        return default


def _report2_period(request):
    today = dj_timezone.localdate()
    filter_type = request.GET.get("filter", "year")
    if filter_type not in REPORT2_PERIOD_TYPES:
        filter_type = "year"

    year_value = today.year
    month_value = today.month
    lot_value = (request.GET.get("lot") or "").strip()
    start_date = today.replace(month=1, day=1)
    end_date = today.replace(month=12, day=31)
    error = ""

    if filter_type == "month":
        year_value = _report2_int(request.GET.get("year"), today.year, 2000, 2100)
        month_value = _report2_int(request.GET.get("month"), today.month, 1, 12)
        try:
            last_day = calendar.monthrange(year_value, month_value)[1]
            start_date = datetime(year_value, month_value, 1).date()
            end_date = datetime(year_value, month_value, last_day).date()
        except ValueError:
            error = "Выбран некорректный месяц."
            start_date = today.replace(month=1, day=1)
            end_date = today.replace(month=12, day=31)
    elif filter_type == "year":
        year_value = _report2_int(request.GET.get("year"), today.year, 2000, 2100)
        try:
            start_date = datetime(year_value, 1, 1).date()
            end_date = datetime(year_value, 12, 31).date()
        except ValueError:
            error = "Выбран некорректный год."
            start_date = today.replace(month=1, day=1)
            end_date = today.replace(month=12, day=31)
    elif filter_type == "lot":
        start_date = None
        end_date = None
    else:
        start_value = request.GET.get("start")
        end_value = request.GET.get("end")
        start_date = _report7_date(start_value)
        end_date = _report7_date(end_value)
        if start_date is None or end_date is None:
            error = "Выбран некорректный период."
            start_date = today.replace(month=1, day=1)
            end_date = today.replace(month=12, day=31)
        elif end_date < start_date:
            error = "Дата начала не может быть позже даты окончания."
            start_date = today.replace(month=1, day=1)
            end_date = today.replace(month=12, day=31)

    return {
        "filter_type": filter_type,
        "start_date": start_date,
        "end_date": end_date,
        "month_value": month_value,
        "year_value": year_value,
        "lot_value": lot_value,
        "error": error,
    }


def _report2_summary(data, start_date, end_date):
    total_days = max(0, (end_date - start_date).days + 1) if start_date and end_date else 0

    fol_hours_total = sum((row.get("fol_hours") or 0) for row in data)
    fol_production_total = sum((row.get("fol_production") or 0) for row in data)
    bol_hours_total = sum((row.get("bol_hours") or 0) for row in data)
    bol_production_total = sum((row.get("bol_production") or 0) for row in data)

    fol_active_days = sum(1 for row in data if (row.get("fol_hours") or 0) > 0)
    bol_active_days = sum(1 for row in data if (row.get("bol_hours") or 0) > 0)

    # Average speed = average of daily speeds where hours > 0
    fol_speeds = [row.get("fol_speed", 0) for row in data if (row.get("fol_hours") or 0) > 0]
    bol_speeds = [row.get("bol_speed", 0) for row in data if (row.get("bol_hours") or 0) > 0]
    fol_speed_avg = round(sum(fol_speeds) / len(fol_speeds), 2) if fol_speeds else 0
    bol_speed_avg = round(sum(bol_speeds) / len(bol_speeds), 2) if bol_speeds else 0

    return {
        "total_days": total_days,
        "fol_hours": round(fol_hours_total, 2),
        "fol_production": fol_production_total,
        "fol_speed": fol_speed_avg,
        "fol_active_days": fol_active_days,
        "bol_hours": round(bol_hours_total, 2),
        "bol_production": bol_production_total,
        "bol_speed": bol_speed_avg,
        "bol_active_days": bol_active_days,
    }


def _report2_data_dates(data):
    """Минимальная и максимальная даты строк данных (для отображения диапазона лота)."""
    dates = [row["report_date"] for row in data if row.get("report_date")]
    if not dates:
        return None, None
    return min(dates), max(dates)


def _report7_summary(data, start_date, end_date):
    lots = sorted({row.get("lot_number") for row in data if row.get("lot_number")})
    unique_pcs = {row.get("pcs_no") for row in data if row.get("pcs_no")}
    return {
        "total_days": max(0, (end_date - start_date).days + 1),
        "released_count": len(unique_pcs),
        "lots": ", ".join(lots),
    }


class HomeView(TemplateView):
    template_name = "home.html"


class ReportListView(TemplateView):
    template_name = "reports/report_list.html"


# ========================
# Отчет 1: Сводный (заказы)
# ========================

REPORT1_PERIOD_TYPES = {"year", "period"}


def _report1_summary(data, start_date, end_date):
    return {
        "total_days": max(0, (end_date - start_date).days + 1),
        "released_lots": sum(1 for row in data if (row.get("Ok") or 0) > 0),
        "production_total": sum(row.get("Production", 0) or 0 for row in data),
        "ok_total": sum(row.get("Ok", 0) or 0 for row in data),
        "diff_total": sum(row.get("Diff", 0) or 0 for row in data),
    }


class Report1View(TemplateView):
    template_name = "reports/report1.html"

    def get(self, request, *args, **kwargs):
        today = dj_timezone.localdate()
        filter_type = request.GET.get("filter", "year")
        if filter_type not in REPORT1_PERIOD_TYPES:
            filter_type = "year"

        year_value = today.year
        start_date = today.replace(month=1, day=1)
        end_date = today.replace(month=12, day=31)
        error = ""

        if filter_type == "year":
            try:
                year_value = int(request.GET.get("year", today.year))
                start_date = datetime(year_value, 1, 1).date()
                end_date = datetime(year_value, 12, 31).date()
            except (ValueError, TypeError):
                error = "Выбран некорректный год."
                year_value = today.year
                start_date = today.replace(month=1, day=1)
                end_date = today.replace(month=12, day=31)
        else:
            start_value = request.GET.get("start")
            end_value = request.GET.get("end")
            start_date = _report7_date(start_value)
            end_date = _report7_date(end_value)
            if start_date is None or end_date is None:
                error = "Выбран некорректный период."
                start_date = today.replace(month=1, day=1)
                end_date = today.replace(month=12, day=31)
            elif end_date < start_date:
                error = "Дата начала не может быть позже даты окончания."
                start_date = today.replace(month=1, day=1)
                end_date = today.replace(month=12, day=31)

        data = report1_orders(start_date, end_date)

        sort_field, sort_dir = get_sort_params(request, "report1")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report1"])
        else:
            # default sort by start ascending
            data = apply_sort(data, "start", "asc", SORTABLE_FIELDS["report1"])
        for i, row in enumerate(data, 1):
            row["N"] = i

        if error:
            messages.error(request, error)

        return self.render_to_response(self.get_context_data(
            data=data,
            summary=_report1_summary(data, start_date, end_date),
            filter_type=filter_type,
            start_date=start_date,
            end_date=end_date,
            year_value=year_value,
            years=range(2024, today.year + 2),
            sort_field=sort_field,
            sort_dir=sort_dir,
        ))


@method_decorator(admin_required, name="dispatch")
class LotInfoUpdateView(UpdateView):
    model = LotInfo
    fields = ["plan_total", "comment"]
    template_name = "reports/lot_edit_modal.html"
    slug_field = "lot_number"
    slug_url_kwarg = "lot_number"

    def get_success_url(self):
        return reverse("reports:report1") + "?edited=1"


# ========================
# Отчет 2: Производительность линии
# ========================

class Report2View(TemplateView):
    template_name = "reports/report2.html"

    def get(self, request, *args, **kwargs):
        period = _report2_period(request)
        data = report2_line_productivity(period["start_date"], period["end_date"], period["lot_value"])

        sort_field, sort_dir = get_sort_params(request, "report2")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report2"])

        # Для фильтра «По лоту» (без диапазона дат) вычисляем диапазон из данных,
        # чтобы сводка и индикаторы «Всего дней» были осмысленными.
        if period["filter_type"] == "lot":
            lot_start, lot_end = _report2_data_dates(data)
            period["start_date"] = lot_start
            period["end_date"] = lot_end

        # Filter to only show days with data
        filtered_data = [row for row in data if row.get("has_data")]
        summary = _report2_summary(filtered_data, period["start_date"], period["end_date"])
        chart_data = {
            "dates": [row["report_date"].isoformat() for row in filtered_data],
            "fol_hours": [float(row.get("fol_hours") or 0) for row in filtered_data],
            "bol_hours": [float(row.get("bol_hours") or 0) for row in filtered_data],
            "fol_production": [float(row.get("fol_production") or 0) for row in filtered_data],
            "bol_production": [float(row.get("bol_production") or 0) for row in filtered_data],
            "fol_speed": [float(row.get("fol_speed") or 0) for row in filtered_data],
            "bol_speed": [float(row.get("bol_speed") or 0) for row in filtered_data],
        }

        if period["error"]:
            messages.error(request, period["error"])

        today = dj_timezone.localdate()
        response = self.render_to_response(self.get_context_data(
            data=filtered_data,
            summary=summary,
            table_summary=summary,
            chart_data=json.dumps(chart_data),
            months=range(1, 13),
            years=range(2024, today.year + 2),
            sort_field=sort_field,
            sort_dir=sort_dir,
            **period,
        ))
        response["Cache-Control"] = "no-store"
        return response


REPORT3_PERIOD_TYPES = {"year"}


def _report3_period(request):
    """Обработчик периода отчёта №3."""
    today = dj_timezone.localdate()
    filter_type = request.GET.get("filter", "year")
    if filter_type not in REPORT3_PERIOD_TYPES:
        filter_type = "year"

    year_value = today.year
    error = ""
    year_param = request.GET.get("year")
    if year_param not in (None, ""):
        try:
            parsed_year = int(year_param)
            if 2000 <= parsed_year <= 2100:
                year_value = parsed_year
            else:
                error = "Выбран некорректный год."
        except (TypeError, ValueError):
            error = "Выбран некорректный год."

    start_date = datetime(year_value, 1, 1).date()
    end_date = datetime(year_value, 12, 31).date()

    return {
        "filter_type": filter_type,
        "year_value": year_value,
        "start_date": start_date,
        "end_date": end_date,
        "error": error,
    }


# ========================
# Отчет 3: Сводный по месяцам
# ========================

class Report3View(TemplateView):
    template_name = "reports/report3.html"

    def get(self, request, *args, **kwargs):
        period = _report3_period(request)
        daily = report2_line_productivity(period["start_date"], period["end_date"])
        data = _aggregate_monthly(daily, period["year_value"])

        sort_field, sort_dir = get_sort_params(request, "report3")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report3"])

        filtered_daily = [row for row in daily if row.get("has_data")]
        summary = _report2_summary(filtered_daily, period["start_date"], period["end_date"])
        table_summary = summary

        chart_data = {
            "months": [row["report_month"].strftime("%Y-%m") for row in data],
            "fol_hours": [float(row.get("fol_hours") or 0) for row in data],
            "bol_hours": [float(row.get("bol_hours") or 0) for row in data],
            "fol_production": [float(row.get("fol_production") or 0) for row in data],
            "bol_production": [float(row.get("bol_production") or 0) for row in data],
            "fol_speed": [float(row.get("fol_speed") or 0) for row in data],
            "bol_speed": [float(row.get("bol_speed") or 0) for row in data],
        }

        today = dj_timezone.localdate()
        response = self.render_to_response(self.get_context_data(
            data=data,
            summary=summary,
            table_summary=table_summary,
            chart_data=json.dumps(chart_data),
            months=range(1, 13),
            years=range(2024, today.year + 2),
            sort_field=sort_field,
            sort_dir=sort_dir,
            **period,
        ))
        response["Cache-Control"] = "no-store"
        return response


# ========================
# Отчет 4: Брак
# ========================

class Report4View(TemplateView):
    template_name = "reports/report4.html"

    def get(self, request, *args, **kwargs):
        filter_type = request.GET.get("filter", "30d")
        lot_filter = request.GET.get("lot", "")

        if filter_type == "period":
            s = request.GET.get("start")
            e = request.GET.get("end")
            if s and e:
                start_date = datetime.strptime(s, "%Y-%m-%d").date()
                end_date = datetime.strptime(e, "%Y-%m-%d").date()
            else:
                start_date, end_date = get_date_range("30d")
        elif filter_type == "lot":
            start_date, end_date = None, None
        else:
            start_date, end_date = get_date_range(filter_type)

        sort_field, sort_dir = get_sort_params(request, "report4")
        limit = 30 if not lot_filter and not sort_field and filter_type == "30d" else None
        data = report4_defects(start_date, end_date, limit=limit)

        if lot_filter:
            data = [d for d in data if lot_filter.lower() in (d.get("lot_number") or "").lower()]

        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report4"])

        return self.render_to_response(self.get_context_data(
            data=data, filter_type=filter_type, start_date=start_date, end_date=end_date,
            lot_filter=lot_filter, sort_field=sort_field, sort_dir=sort_dir,
        ))


@method_decorator(admin_required, name="dispatch")
class ManualDefectCreateView(CreateView):
    model = ManualDefect
    fields = ["pcs_no", "production_spec", "lot_number", "entry_date", "reason", "comment"]
    template_name = "reports/defect_form.html"

    def get_success_url(self):
        return reverse("reports:report4") + "?added=1"

    def form_valid(self, form):
        form.instance.created_by = getattr(self.request.user, "username", "admin")
        return super().form_valid(form)


class SNCommentForm(forms.ModelForm):
    class Meta:
        model = SNComment
        fields = ["comment"]


@method_decorator(admin_required, name="dispatch")
class SNCommentEditView(View):
    template_name = "reports/sn_comment_form.html"

    def get(self, request, pcs_no):
        if not self._pcs_exists(pcs_no):
            raise Http404("Серийный номер не найден.")
        comment = SNComment.objects.filter(pcs_no=pcs_no).first()
        if comment is None:
            comment = SNComment(pcs_no=pcs_no)
        return render(
            request,
            self.template_name,
            {"form": SNCommentForm(instance=comment), "pcs_no": pcs_no},
        )

    def post(self, request, pcs_no):
        if not self._pcs_exists(pcs_no):
            raise Http404("Серийный номер не найден.")
        comment = SNComment.objects.filter(pcs_no=pcs_no).first()
        if comment is None:
            comment = SNComment(pcs_no=pcs_no)
        form = SNCommentForm(request.POST, instance=comment)
        if form.is_valid():
            comment = form.save(commit=False)
            comment.pcs_no = pcs_no
            if not comment.created_by:
                comment.created_by = request.user.get_username()
            comment.save()
            messages.success(request, "Комментарий сохранен.")
            return redirect("reports:report4")
        return render(
            request,
            self.template_name,
            {"form": form, "pcs_no": pcs_no},
            status=400,
        )

    @staticmethod
    def _pcs_exists(pcs_no):
        return (
            bool(pcs_no)
            and len(pcs_no) <= SNComment._meta.get_field("pcs_no").max_length
            and (
                ProductionRecord.objects.filter(pcs_no=pcs_no).exists()
                or ManualDefect.objects.filter(pcs_no=pcs_no).exists()
            )
        )


# ========================
# Отчет 5: Повторные проходы
# ========================

class Report5View(TemplateView):
    template_name = "reports/report5.html"

    def get(self, request, *args, **kwargs):
        filter_type = request.GET.get("filter", "30d")
        lot_filter = request.GET.get("lot", "")

        if filter_type == "period":
            s = request.GET.get("start")
            e = request.GET.get("end")
            if s and e:
                start_date = datetime.strptime(s, "%Y-%m-%d").date()
                end_date = datetime.strptime(e, "%Y-%m-%d").date()
            else:
                start_date, end_date = get_date_range("30d")
        elif filter_type == "lot":
            start_date, end_date = None, None
        else:
            start_date, end_date = get_date_range(filter_type)

        sort_field, sort_dir = get_sort_params(request, "report5")
        limit = 50 if not lot_filter and not start_date and not sort_field else None
        data = report5_repeated_passes(start_date, end_date, limit=limit)

        if lot_filter:
            data = [d for d in data if lot_filter.lower() in (d.get("lot_number") or "").lower()]

        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report5"])

        return self.render_to_response(self.get_context_data(
            data=data, filter_type=filter_type, start_date=start_date, end_date=end_date,
            lot_filter=lot_filter, sort_field=sort_field, sort_dir=sort_dir,
        ))


@method_decorator(admin_required, name="dispatch")
class Report5CommentFormView(View):
    template_name = "reports/report5_comment_form.html"

    def get(self, request):
        return render(request, self.template_name)

    def post(self, request):
        pcs_no = request.POST.get("pcs_no", "")
        station_no = request.POST.get("station_no")
        comment = request.POST.get("comment", "")
        if pcs_no and station_no:
            try:
                station_no = int(station_no)
                Report5Comment.objects.update_or_create(
                    pcs_no=pcs_no, station_no=station_no,
                    defaults={"comment": comment, "created_by": getattr(request.user, "username", "admin")},
                )
                messages.success(request, "Комментарий сохранен.")
            except ValueError:
                messages.error(request, "Введите корректный номер станции.")
        return redirect("reports:report5")


# ========================
# Отчет 6: Серийный номер
# ========================

class Report6View(TemplateView):
    template_name = "reports/report6.html"

    def get(self, request, *args, **kwargs):
        pcs_no = request.GET.get("pcs_no", "").strip()
        data = report6_serial_number(pcs_no) if pcs_no else []
        exists = ProductionRecord.objects.filter(pcs_no=pcs_no).exists() if pcs_no else None

        if pcs_no and not exists:
            messages.error(request, f'Серийный номер "{pcs_no}" не найден в базе.')

        sort_field, sort_dir = get_sort_params(request, "report6")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report6"])

        return self.render_to_response(self.get_context_data(
            data=data, pcs_no=pcs_no, exists=exists,
            sort_field=sort_field, sort_dir=sort_dir,
        ))


class Report7View(TemplateView):
    template_name = "reports/report7.html"

    def get(self, request, *args, **kwargs):
        period = _report7_period(request)
        data = report7_station130_output(period["start_date"], period["end_date"])

        sort_field, sort_dir = get_sort_params(request, "report7")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report7"])

        if period["error"]:
            messages.error(request, period["error"])

        return self.render_to_response(self.get_context_data(
            data=data,
            summary=_report7_summary(data, period["start_date"], period["end_date"]),
            months=range(1, 13),
            years=range(2024, dj_timezone.localdate().year + 2),
            sort_field=sort_field,
            sort_dir=sort_dir,
            **period,
        ))


# ========================
# Экспорт в Excel
# ========================

class _ExcelExportBase(View):
    """Базовый класс для экспорта в Excel."""

    @staticmethod
    def _excel_value(value):
        """Remove tzinfo from timezone-aware datetimes (Excel unsupported)."""
        if isinstance(value, datetime) and value.tzinfo is not None:
            return dj_timezone.make_naive(value) if value.tzinfo is not None else value
        return value

    def _make_response(self, wb, filename):
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class Report1ExportView(_ExcelExportBase):
    def get(self, request):
        filter_type = request.GET.get("filter", "30d")
        if filter_type == "period":
            s = request.GET.get("start")
            e = request.GET.get("end")
            if s and e:
                start_date = datetime.strptime(s, "%Y-%m-%d").date()
                end_date = datetime.strptime(e, "%Y-%m-%d").date()
            else:
                start_date, end_date = get_date_range("30d")
        else:
            start_date, end_date = get_date_range(filter_type)
        data = report1_orders(start_date, end_date)

        sort_field, sort_dir = get_sort_params(request, "report1")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report1"])
            for i, row in enumerate(data, 1):
                row["N"] = i

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Сводный отчет"
        headers = ["N", "LOT", "Production spec.", "Total", "Production", "Ok", "Diff", "Defect", "start", "finish", "comment"]
        for col, h in enumerate(headers, 1):
            ws.cell(1, col, h)
        for row_num, d in enumerate(data, 2):
            ws.cell(row_num, 1, d["N"])
            ws.cell(row_num, 2, d["LOT"])
            ws.cell(row_num, 3, d["production_spec"])
            ws.cell(row_num, 4, d["Total"])
            ws.cell(row_num, 5, d["Production"])
            ws.cell(row_num, 6, d["Ok"])
            ws.cell(row_num, 7, d["Diff"])
            ws.cell(row_num, 8, d["Defect"])
            ws.cell(row_num, 9, self._excel_value(d["start"]))
            ws.cell(row_num, 10, self._excel_value(d["finish"]))
            ws.cell(row_num, 11, d["comment"])

        return self._make_response(wb, "report1_orders.xlsx")


class Report2ExportView(_ExcelExportBase):
    fol_color = "4361EE"
    bol_color = "EF4444"
    fol_fill = "EAF6FF"
    bol_fill = "FFF0F0"
    header_fill = "1F4E78"

    @staticmethod
    def _style_report2_header(cell, fill_color):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="center")

    @staticmethod
    def _write_report2_summary(ws, period, summary, title="Отчёт №2: Производительность сборочной линии"):
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        ws["A1"].alignment = Alignment(horizontal="center")

        summary_rows = [
            ["Период", f'{period["start_date"]:%d.%m.%Y} — {period["end_date"]:%d.%m.%Y}'],
            ["Дата начала", period["start_date"]],
            ["Дата окончания", period["end_date"]],
            ["Всего дней", summary["total_days"]],
            [],
            ["Показатель", "FOL", "BOL"],
            ["Часы работы", summary["fol_hours"], summary["bol_hours"]],
            ["Выпуск продукции", summary["fol_production"], summary["bol_production"]],
            ["Средняя скорость, шт/ч", summary["fol_speed"], summary["bol_speed"]],
            ["Активных дней", summary["fol_active_days"], summary["bol_active_days"]],
        ]
        for row_num, row in enumerate(summary_rows, 2):
            for col_num, value in enumerate(row, 1):
                ws.cell(row_num, col_num, value)

        for cell in ws[7]:
            fill_color = "1F4E78"
            if cell.column == 2:
                fill_color = Report2ExportView.fol_color
            elif cell.column == 3:
                fill_color = Report2ExportView.bol_color
            Report2ExportView._style_report2_header(cell, fill_color)
        for row_num in range(8, 12):
            ws.cell(row_num, 2).fill = PatternFill("solid", fgColor=Report2ExportView.fol_fill)
            ws.cell(row_num, 3).fill = PatternFill("solid", fgColor=Report2ExportView.bol_fill)
        for row_num in range(2, 6):
            ws.cell(row_num, 1).font = Font(bold=True)
        for row_num in range(3, 5):
            ws.cell(row_num, 2).number_format = "dd.mm.yyyy"
        for row_num in range(8, 11):
            ws.cell(row_num, 2).number_format = "0.00"
            ws.cell(row_num, 3).number_format = "0.00"

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18
        ws.freeze_panes = "A7"

    @staticmethod
    def _add_report2_chart(chart_ws, data_ws, title, fol_col, bol_col, anchor):
        last_row = data_ws.max_row
        if last_row < 2:
            return

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = title
        chart.y_axis.title = "Значение"
        chart.x_axis.title = "Дата"
        chart.height = 8
        chart.width = 15
        chart.gapWidth = 50

        fol_data = Reference(data_ws, min_col=fol_col, min_row=1, max_row=last_row)
        bol_data = Reference(data_ws, min_col=bol_col, min_row=1, max_row=last_row)
        categories = Reference(data_ws, min_col=1, min_row=2, max_row=last_row)
        chart.add_data(fol_data, titles_from_data=True)
        chart.add_data(bol_data, titles_from_data=True)
        chart.set_categories(categories)
        chart.series[0].graphicalProperties.solidFill = Report2ExportView.fol_color
        chart.series[0].graphicalProperties.line.solidFill = Report2ExportView.fol_color
        chart.series[1].graphicalProperties.solidFill = Report2ExportView.bol_color
        chart.series[1].graphicalProperties.line.solidFill = Report2ExportView.bol_color
        chart.legend.position = "b"
        chart_ws.add_chart(chart, anchor)

    def get(self, request):
        period = _report2_period(request)
        data = report2_line_productivity(period["start_date"], period["end_date"], period["lot_value"])

        sort_field, sort_dir = get_sort_params(request, "report2")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report2"])

        # Для фильтра «По лоту» вычисляем диапазон из данных для сводки.
        if period["filter_type"] == "lot":
            lot_start, lot_end = _report2_data_dates(data)
            period["start_date"] = lot_start
            period["end_date"] = lot_end

        data = [row for row in data if row.get("has_data")]
        summary = _report2_summary(data, period["start_date"], period["end_date"])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Производительность"
        headers = ["Дата", "FOL hours", "FOL production", "FOL speed", "BOL hours", "BOL production", "BOL speed"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            fill_color = self.header_fill
            if 2 <= col <= 4:
                fill_color = self.fol_color
            elif col >= 5:
                fill_color = self.bol_color
            self._style_report2_header(cell, fill_color)

        for row_num, row in enumerate(data, 2):
            ws.cell(row_num, 1, row["report_date"].strftime("%Y-%m-%d"))
            ws.cell(row_num, 2, row["fol_hours"])
            ws.cell(row_num, 3, row["fol_production"])
            ws.cell(row_num, 4, row["fol_speed"])
            ws.cell(row_num, 5, row["bol_hours"])
            ws.cell(row_num, 6, row["bol_production"])
            ws.cell(row_num, 7, row["bol_speed"])
            for col in range(2, 5):
                ws.cell(row_num, col).fill = PatternFill("solid", fgColor=self.fol_fill)
            for col in range(5, 8):
                ws.cell(row_num, col).fill = PatternFill("solid", fgColor=self.bol_fill)
            for col in range(2, 8):
                ws.cell(row_num, col).number_format = "0.00"

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:G{ws.max_row}"
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 14

        summary_ws = wb.create_sheet("Сводка")
        self._write_report2_summary(summary_ws, period, summary)

        charts_ws = wb.create_sheet("Диаграммы")
        charts_ws.sheet_view.showGridLines = False
        self._add_report2_chart(charts_ws, ws, "Время работы, ч", 2, 5, "A3")
        self._add_report2_chart(charts_ws, ws, "Выпуск продукции, шт", 3, 6, "J3")
        self._add_report2_chart(charts_ws, ws, "Скорость выпуска, шт/ч", 4, 7, "A20")

        return self._make_response(wb, "report2_line.xlsx")


class Report3ExportView(Report2ExportView):
    """Экспорт отчёта №3: ежемесячная сводка, сводка и диаграммы."""

    def get(self, request):
        period = _report3_period(request)
        daily = report2_line_productivity(period["start_date"], period["end_date"])
        data = _aggregate_monthly(daily, period["year_value"])

        sort_field, sort_dir = get_sort_params(request, "report3")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report3"])

        filtered_daily = [row for row in daily if row.get("has_data")]
        summary = _report2_summary(filtered_daily, period["start_date"], period["end_date"])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Сводный месяцы"
        headers = ["Месяц", "FOL hours", "FOL production", "FOL speed",
                   "BOL hours", "BOL production", "BOL speed"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            fill_color = self.header_fill
            if 2 <= col <= 4:
                fill_color = self.fol_color
            elif col >= 5:
                fill_color = self.bol_color
            self._style_report2_header(cell, fill_color)

        for row_num, row in enumerate(data, 2):
            ws.cell(row_num, 1, row["report_month"].strftime("%Y-%m"))
            ws.cell(row_num, 2, row["fol_hours"])
            ws.cell(row_num, 3, row["fol_production"])
            ws.cell(row_num, 4, row["fol_speed"])
            ws.cell(row_num, 5, row["bol_hours"])
            ws.cell(row_num, 6, row["bol_production"])
            ws.cell(row_num, 7, row["bol_speed"])
            for col in range(2, 5):
                ws.cell(row_num, col).fill = PatternFill("solid", fgColor=self.fol_fill)
            for col in range(5, 8):
                ws.cell(row_num, col).fill = PatternFill("solid", fgColor=self.bol_fill)
            for col in range(2, 8):
                ws.cell(row_num, col).number_format = "0.00"

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:G{ws.max_row}"
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 14

        summary_ws = wb.create_sheet("Сводка")
        self._write_report2_summary(
            summary_ws,
            period,
            summary,
            title="Отчёт №3: Сводный по месяцам",
        )

        charts_ws = wb.create_sheet("Диаграммы")
        charts_ws.sheet_view.showGridLines = False
        self._add_report2_chart(charts_ws, ws, "Время работы, ч", 2, 5, "A3")
        self._add_report2_chart(charts_ws, ws, "Выпуск продукции, шт", 3, 6, "J3")
        self._add_report2_chart(charts_ws, ws, "Скорость выпуска, шт/ч", 4, 7, "A20")

        response = self._make_response(wb, "report3_monthly.xlsx")
        response["Cache-Control"] = "no-store"
        return response


class Report4ExportView(_ExcelExportBase):
    def get(self, request):
        filter_type = request.GET.get("filter", "30d")
        lot_filter = request.GET.get("lot", "")

        if filter_type == "period":
            s = request.GET.get("start")
            e = request.GET.get("end")
            if s and e:
                start_date = datetime.strptime(s, "%Y-%m-%d").date()
                end_date = datetime.strptime(e, "%Y-%m-%d").date()
            else:
                start_date, end_date = get_date_range("30d")
        elif filter_type == "lot":
            start_date, end_date = None, None
        else:
            start_date, end_date = get_date_range(filter_type)

        sort_field, sort_dir = get_sort_params(request, "report4")
        limit = 30 if not lot_filter and not sort_field and filter_type == "30d" else None
        data = report4_defects(start_date, end_date, limit=limit)
        if lot_filter:
            data = [d for d in data if lot_filter.lower() in (d.get("lot_number") or "").lower()]

        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report4"])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Брак"
        headers = [
            "№", "Серийный номер", "Production spec.", "Лот",
            "Дата поступления в производство", "Последняя станция",
            "Дата прохождения последней станции", "Комментарий",
        ]
        for col, h in enumerate(headers, 1):
            ws.cell(1, col, h)
        for row_num, d in enumerate(data, 2):
            ws.cell(row_num, 1, row_num - 1)
            ws.cell(row_num, 2, d["pcs_no"])
            ws.cell(row_num, 3, d["production_spec"])
            ws.cell(row_num, 4, d["lot_number"] or "")
            ws.cell(row_num, 5, self._excel_value(d["entry_date"]))
            ws.cell(row_num, 6, d.get("last_station", ""))
            ws.cell(row_num, 7, self._excel_value(d.get("last_station_date")))
            ws.cell(row_num, 8, d.get("comment", ""))

        return self._make_response(wb, "report4_defects.xlsx")


class Report5ExportView(_ExcelExportBase):
    def get(self, request):
        filter_type = request.GET.get("filter", "30d")
        lot_filter = request.GET.get("lot", "")

        if filter_type == "period":
            s = request.GET.get("start")
            e = request.GET.get("end")
            if s and e:
                start_date = datetime.strptime(s, "%Y-%m-%d").date()
                end_date = datetime.strptime(e, "%Y-%m-%d").date()
            else:
                start_date, end_date = get_date_range("30d")
        elif filter_type == "lot":
            start_date, end_date = None, None
        else:
            start_date, end_date = get_date_range(filter_type)

        sort_field, sort_dir = get_sort_params(request, "report5")
        limit = 50 if not lot_filter and not start_date and not sort_field else None
        data = report5_repeated_passes(start_date, end_date, limit=limit)
        if lot_filter:
            data = [d for d in data if lot_filter.lower() in (d.get("lot_number") or "").lower()]

        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report5"])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Повторные проходы"
        headers = ["№", "Серийный номер", "Production spec.", "Лот", "Станция", "Дата время", "Комментарий"]
        for col, h in enumerate(headers, 1):
            ws.cell(1, col, h)
        for row_num, d in enumerate(data, 2):
            ws.cell(row_num, 1, row_num - 1)
            ws.cell(row_num, 2, d["pcs_no"])
            ws.cell(row_num, 3, d["production_spec"])
            ws.cell(row_num, 4, d["lot_number"] or "")
            ws.cell(row_num, 5, d["station"])
            ws.cell(row_num, 6, d["dates"])
            ws.cell(row_num, 7, d.get("comment", ""))

        return self._make_response(wb, "report5_repeats.xlsx")


class Report6ExportView(_ExcelExportBase):
    def get(self, request):
        pcs_no = request.GET.get("pcs_no", "").strip()
        data = report6_serial_number(pcs_no) if pcs_no else []
        if pcs_no and not data:
            return redirect("reports:report6")

        sort_field, sort_dir = get_sort_params(request, "report6")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report6"])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Отчет по SN"
        headers = [
            "PCSNo", "Lot no.", "Production spec.", "Subop no",
            "WorkstationName", "CREATEDATE", "result", "testData", "Комментарий",
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(1, col, header)
        for row_num, d in enumerate(data, 2):
            ws.cell(row_num, 1, d["pcs_no"])
            ws.cell(row_num, 2, d["lot_number"])
            ws.cell(row_num, 3, d["production_spec"])
            ws.cell(row_num, 4, d["subop_no"])
            ws.cell(row_num, 5, d["workstation_name"])
            ws.cell(row_num, 6, self._excel_value(d["created_date"]))
            ws.cell(row_num, 7, d["result"])
            ws.cell(row_num, 8, d["test_data"])
            ws.cell(row_num, 9, d.get("comment", ""))

        return self._make_response(wb, f"report6_{pcs_no}.xlsx")


class Report7ExportView(_ExcelExportBase):
    def get(self, request):
        period = _report7_period(request)
        data = report7_station130_output(period["start_date"], period["end_date"])

        sort_field, sort_dir = get_sort_params(request, "report7")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report7"])

        summary = _report7_summary(data, period["start_date"], period["end_date"])
        wb = openpyxl.Workbook()
        summary_ws = wb.active
        summary_ws.title = "Сводка"
        summary_rows = [
            ["Период", f'{period["start_date"].isoformat()} — {period["end_date"].isoformat()}'],
            ["Всего дней", summary["total_days"]],
            ["Выпущено продукции", summary["released_count"]],
            ["Лоты", summary["lots"]],
        ]
        for row_num, row in enumerate(summary_rows, 1):
            summary_ws.cell(row_num, 1, row[0])
            summary_ws.cell(row_num, 2, row[1])

        ws = wb.create_sheet("Выпуск")
        headers = [
            "PCSNo - Серийный номер",
            "Lot no. - Лот",
            "Production spec. - Спецификация",
            "Date - Дата прохождения",
            "Time - Время прохождения",
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(1, col, header)
        for row_num, d in enumerate(data, 2):
            created_date = d["created_date"]
            ws.cell(row_num, 1, d["pcs_no"])
            ws.cell(row_num, 2, d["lot_number"])
            ws.cell(row_num, 3, d["production_spec"])
            ws.cell(row_num, 4, created_date.date())
            ws.cell(row_num, 5, created_date.time().replace(tzinfo=None))

        filename = (
            f'report7_{period["start_date"].isoformat()}'
            f'_{period["end_date"].isoformat()}.xlsx'
        )
        return self._make_response(wb, filename)


# ========================
# Отчет 8: По станциям
# ========================

REPORT8_PERIOD_TYPES = {"year", "month", "period", "lot"}

REPORT8_CHART_LABELS = [label for _subop, label in REPORT8_CHART_STATIONS]


REPORT8_DESCRIPTION = """\
Отчёт «По станциям» показывает дневное качество прохождения каждой станции сборочной линии.

Для каждой даты (в которую был зафиксирован ход производства) отображаются:
  • уникальное количество серийных номеров (SN), прошедших станцию с результатом OK — для промежуточных станций Подача, Плазма, Сборка, Стабилизация, Burn-in;
  • количество всех записей OK и NG на каждой тестовой станции (FCT1, FCT2, Прошивка, Проверка прошивки, Герметичность, PIN-тест, Этикетка);
  • процент NG от OK: 0 % выделяются зелёным, больше 0 % — красным;
  • список серийных номеров с результатом NG — каждый номер — ссылка на Отчёт по SN;
  • разница FOL = (уникальные OK на 010) − (уникальные OK на 050);
  • разница BOL = (уникальные OK на 060) − (уникальные OK на 130).
Разница окрашивается: 0 — зелёный, < 0 — красный, > 0 — светло-фиолетовый.

Фильтры: Год (по умолчанию — последний календарный год), Месяц, Произвольный период, По лоту.
Экспорт в Excel доступен без авторизации.
"""


def _report8_int(value, default, lower, upper):
    try:
        parsed = int(value)
        return parsed if lower <= parsed <= upper else default
    except (TypeError, ValueError):
        return default


def _report8_period(request):
    """Разбор фильтра периода для отчёта 8: year, month, period, lot."""
    today = dj_timezone.localdate()
    filter_type = request.GET.get("filter", "year")
    if filter_type not in REPORT8_PERIOD_TYPES:
        filter_type = "year"

    # По умолчанию — последний календарный год.
    last_year = today.year - 1
    year_value = last_year
    month_value = today.month
    lot_value = (request.GET.get("lot") or "").strip()
    start_date = datetime(last_year, 1, 1).date()
    end_date = datetime(last_year, 12, 31).date()
    error = ""

    if filter_type == "month":
        year_value = _report8_int(request.GET.get("year"), last_year, 2000, 2100)
        month_value = _report8_int(request.GET.get("month"), today.month, 1, 12)
        try:
            last_day = calendar.monthrange(year_value, month_value)[1]
            start_date = datetime(year_value, month_value, 1).date()
            end_date = datetime(year_value, month_value, last_day).date()
        except ValueError:
            error = "Выбран некорректный месяц."
            start_date = datetime(last_year, 1, 1).date()
            end_date = datetime(last_year, 12, 31).date()
            year_value = last_year
            month_value = today.month
    elif filter_type == "year":
        year_value = _report8_int(request.GET.get("year"), last_year, 2000, 2100)
        try:
            start_date = datetime(year_value, 1, 1).date()
            end_date = datetime(year_value, 12, 31).date()
        except ValueError:
            error = "Выбран некорректный год."
            year_value = last_year
            start_date = datetime(last_year, 1, 1).date()
            end_date = datetime(last_year, 12, 31).date()
    elif filter_type == "period":
        start_value = request.GET.get("start")
        end_value = request.GET.get("end")
        parsed_start = _report7_date(start_value)
        parsed_end = _report7_date(end_value)
        if parsed_start is None or parsed_end is None:
            error = "Выбран некорректный период."
        elif parsed_end < parsed_start:
            error = "Дата начала не может быть позже даты окончания."
        if error:
            start_date = datetime(last_year, 1, 1).date()
            end_date = datetime(last_year, 12, 31).date()
        else:
            start_date = parsed_start
            end_date = parsed_end
    else:  # lot
        start_date = None
        end_date = None

    return {
        "filter_type": filter_type,
        "start_date": start_date,
        "end_date": end_date,
        "year_value": year_value,
        "month_value": month_value,
        "lot_value": lot_value,
        "error": error,
    }


def _report8_summary(data):
    """Сводка: суммы для счётчиков, средние для %, суммы для разностей."""
    # Карта station_no -> (ok_key, ng_key) для столбцов %.
    station_count_keys = {}
    for key, station_no, kind, _r1, _r2, _css in REPORT8_COLUMNS:
        if kind == "all_ok":
            station_count_keys.setdefault(station_no, {})["ok"] = key
        elif kind == "all_ng":
            station_count_keys.setdefault(station_no, {})["ng"] = key

    summary = {}
    for key, station_no, kind, _r1, _r2, _css in REPORT8_COLUMNS:
        if kind == "date":
            summary[key] = ""
            continue
        if kind == "ng_sns":
            summary[key] = []
            continue
        if kind in ("unique_ok", "all_ok", "all_ng", "diff_fol", "diff_bol"):
            summary[key] = sum((row.get(key) or 0) for row in data)
        elif kind == "pct":
            count_keys = station_count_keys.get(station_no, {})
            ok_key = count_keys.get("ok")
            ng_key = count_keys.get("ng")
            values = [
                row[key] for row in data
                if (row.get(ok_key) or 0) + (row.get(ng_key) or 0) > 0
            ]
            summary[key] = round(sum(values) / len(values), 2) if values else 0.0
    return summary


def _build_report8_chart_data(data):
    """Дневные данные для диаграммы «Средний % NG по дням».

    Возвращает список дат (ось X) и набор серий — по одной линии на каждую
    тестовую станцию: FCT1, FCT2, Прошивка, Проверка прошивки, Герметичность,
    PIN, Этикетка. Каждая серия содержит значения % NG для каждой даты,
    каждый NG-тип окрашен в свой цвет.
    """
    dates = [
        (row["report_date"].isoformat() if row.get("report_date") else "")
        for row in data
    ]
    datasets = []
    for (_subop, label), key, color in zip(
        REPORT8_CHART_STATIONS, REPORT8_CHART_KEYS, REPORT8_CHART_COLORS
    ):
        datasets.append({
            "label": label,
            "color": color,
            "data": [float(row.get(key) or 0) for row in data],
        })
    return {"dates": dates, "datasets": datasets}


class Report8View(TemplateView):
    template_name = "reports/report8.html"

    def get(self, request, *args, **kwargs):
        period = _report8_period(request)
        data = report8_stations(period["start_date"], period["end_date"], period["lot_value"])

        sort_field, sort_dir = get_sort_params(request, "report8")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report8"])
        else:
            data = apply_sort(data, "report_date", "asc", SORTABLE_FIELDS["report8"])

        summary = _report8_summary(data)
        chart_data = _build_report8_chart_data(data)
        chart_summary = [
            (label, summary[key])
            for (_subop, label), key in zip(REPORT8_CHART_STATIONS, REPORT8_CHART_KEYS)
        ]

        if period["error"]:
            messages.error(request, period["error"])

        today = dj_timezone.localdate()
        response = self.render_to_response(self.get_context_data(
            data=data,
            summary=summary,
            table_summary=summary,
            chart_data=json.dumps(chart_data),
            chart_summary=chart_summary,
            columns=report8_table_columns(),
            header_groups=REPORT8_HEADER_GROUPS,
            report8_description_lines=REPORT8_DESCRIPTION.strip().split("\n"),
            years=range(2024, today.year + 2),
            months=range(1, 13),
            sort_field=sort_field,
            sort_dir=sort_dir,
            **period,
        ))
        response["Cache-Control"] = "no-store"
        return response


class Report8ExportView(_ExcelExportBase):
    def get(self, request):
        period = _report8_period(request)
        data = report8_stations(period["start_date"], period["end_date"], period["lot_value"])

        sort_field, sort_dir = get_sort_params(request, "report8")
        if sort_field:
            data = apply_sort(data, sort_field, sort_dir, SORTABLE_FIELDS["report8"])
        else:
            data = apply_sort(data, "report_date", "asc", SORTABLE_FIELDS["report8"])

        summary = _report8_summary(data)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "По станциям"

        green_fill = "D4EDDA"
        red_fill = "F8D7DA"
        violet_fill = "E6E6FA"

        # Row 1: station titles with merged cells for groups.
        col = 1
        group_col_start = {}
        for group in REPORT8_HEADER_GROUPS:
            title = group["title"]
            span = group["span"]
            if span > 1:
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + span - 1)
            cell = ws.cell(1, col, title)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if span > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center", textRotation=0)
            group_col_start[col] = title
            col += span

        # Row 2: sub-headers.
        for idx, (key, station_no, kind, _r1, r2_title, css) in enumerate(REPORT8_COLUMNS, 1):
            cell = ws.cell(2, idx, r2_title or "")
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(bold=True, color="FFFFFF")

        data_ws = ws
        for row_num, row in enumerate(data, 3):
            for idx, (key, station_no, kind, _r1, _r2, css) in enumerate(REPORT8_COLUMNS, 1):
                cell = data_ws.cell(row_num, idx)
                value = row.get(key)
                if kind == "date":
                    cell.value = self._excel_value(value).strftime("%Y-%m-%d") if value else ""
                elif kind == "ng_sns":
                    cell.value = ", ".join(value) if value else ""
                elif kind == "pct":
                    cell.value = float(value) if value is not None else 0
                    cell.fill = PatternFill("solid", fgColor=green_fill if value == 0 else red_fill)
                    cell.number_format = "0.00"
                elif kind in ("diff_fol", "diff_bol"):
                    cell.value = int(value) if value is not None else 0
                    v = value or 0
                    cell.fill = PatternFill("solid", fgColor=green_fill if v == 0 else (red_fill if v < 0 else violet_fill))
                else:
                    cell.value = int(value) if value is not None else 0

        # Summary row.
        summary_row = len(data) + 3
        data_ws.cell(summary_row, 1, "Итог")
        for idx, (key, station_no, kind, _r1, _r2, css) in enumerate(REPORT8_COLUMNS, 1):
            if idx == 1:
                continue
            cell = data_ws.cell(summary_row, idx)
            cell.font = Font(bold=True)
            value = summary.get(key)
            if kind == "ng_sns":
                cell.value = ""
            elif kind == "pct":
                cell.value = float(value) if value is not None else 0
                cell.fill = PatternFill("solid", fgColor=green_fill if value == 0 else red_fill)
                cell.number_format = "0.00"
            elif kind in ("diff_fol", "diff_bol"):
                cell.value = int(value) if value is not None else 0
                v = value or 0
                cell.fill = PatternFill("solid", fgColor=green_fill if v == 0 else (red_fill if v < 0 else violet_fill))
            else:
                cell.value = int(value) if value is not None else 0

        data_ws.freeze_panes = "A3"
        data_ws.auto_filter.ref = f"A1:{data_ws.cell(2, len(REPORT8_COLUMNS)).coordinate}"

        # Summary sheet.
        summary_ws = wb.create_sheet("Сводка")
        summary_ws["A1"] = "Отчёт 8: По станциям"
        summary_ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        summary_ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        summary_ws["A1"].alignment = Alignment(horizontal="center")

        date_str = ""
        if period["start_date"] and period["end_date"]:
            date_str = f'{period["start_date"].isoformat()} — {period["end_date"].isoformat()}'
        else:
            date_str = "По лоту: " + period["lot_value"] if period["lot_value"] else "без фильтра по дате"

        summary_rows = [
            ["Период", date_str],
            ["Фильтр", period["filter_type"]],
        ]
        for sr, row in enumerate(summary_rows, 2):
            for sc, val in enumerate(row, 1):
                summary_ws.cell(sr, sc, val)

        chart_row = 4
        for idx, (label, key) in enumerate(zip(REPORT8_CHART_LABELS, REPORT8_CHART_KEYS), chart_row):
            summary_ws.cell(idx, 1, label)
            summary_ws.cell(idx, 2, float(summary.get(key) or 0))

        # Дневная диаграмма «Средний % NG по дням»: одна линия на каждую
        # тестовую станцию (FCT1, FCT2, …) — динамика % NG по датам.
        chart_series = _build_report8_chart_data(data)
        chart_dates = chart_series["dates"]
        chart_datasets = chart_series["datasets"]
        n_series = len(chart_datasets)

        if chart_dates:
            daily_header_row = chart_row + len(REPORT8_CHART_LABELS) + 1  # строка 12 (пропуск строки 11)
            # Заголовки колонок дневного графика.
            summary_ws.cell(daily_header_row, 1, "Дата")
            summary_ws.cell(daily_header_row, 1).font = Font(bold=True)
            for col_idx, ds in enumerate(chart_datasets, 2):
                summary_ws.cell(daily_header_row, col_idx, ds["label"])
                summary_ws.cell(daily_header_row, col_idx).font = Font(bold=True)

            # Значения по датам.
            first_data_row = daily_header_row + 1
            last_data_row = daily_header_row + len(chart_dates)
            for row_idx, date_str in enumerate(chart_dates, first_data_row):
                summary_ws.cell(row_idx, 1, date_str)
            for col_idx, ds in enumerate(chart_datasets, 2):
                for row_idx, value in enumerate(ds["data"], first_data_row):
                    summary_ws.cell(row_idx, col_idx, float(value))

            # Линия: одна серия на станцию, каждая своего цвета.
            data_ref = Reference(
                summary_ws, min_col=2, min_row=daily_header_row,
                max_col=2 + n_series - 1, max_row=last_data_row,
            )
            cats_ref = Reference(
                summary_ws, min_col=1, min_row=first_data_row, max_row=last_data_row,
            )
            line_chart = LineChart()
            line_chart.style = 12
            line_chart.title = "Средний % NG по дням"
            line_chart.y_axis.title = "Средний % NG"
            line_chart.x_axis.title = "Дата"
            line_chart.height = 9
            line_chart.width = 22
            line_chart.legend.position = "r"
            line_chart.add_data(data_ref, titles_from_data=True)
            line_chart.set_categories(cats_ref)
            for series, ds in zip(line_chart.series, chart_datasets):
                color_hex = ds["color"].lstrip("#").upper()
                series.graphicalProperties.line.solidFill = color_hex
                series.graphicalProperties.line.width = 2.25
                series.marker = Marker(symbol="circle", size=5)
            summary_ws.add_chart(line_chart, "J12")

        filename = (
            f'report8_stations_'
            f'{period["start_date"].isoformat() if period.get("start_date") else "lot"}_'
            f'{period["end_date"].isoformat() if period.get("end_date") else period["lot_value"]}.xlsx'
        )
        return self._make_response(wb, filename)
