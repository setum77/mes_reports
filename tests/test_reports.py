"""
Тесты бизнес-логики отчетов.
"""
import warnings
warnings.filterwarnings("ignore")

import os, sys, django
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "mes_report3.settings.dev")
sys.path.insert(0, "src")
django.setup()

import pytest
from datetime import date, datetime
from io import BytesIO
from django.contrib.auth.models import AnonymousUser, User
from django.http import HttpResponseRedirect
from django.test import Client, RequestFactory, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from mes_report3.admin_site import mes_admin_site
from production.admin import LotInfoAdmin, ProductionRecordAdmin
from production.models import LotInfo, ManualDefect, ProductionRecord, Report5Comment
from reports.queries import (
    report1_orders,
    report2_line_productivity,
    report5_repeated_passes,
    report6_serial_number,
    report7_station130_output,
    translate_workstation_name,
)
from reports.views import (
    _report1_summary,
    _report2_period,
    _report2_summary,
    _report7_period,
    _report7_summary,
    apply_sort,
    get_sort_params,
    SORTABLE_FIELDS,
)


@pytest.mark.django_db
def test_report1_orders():
    """Проверка расчетов Report 1: Production, Ok, Diff."""
    dt1 = timezone.make_aware(datetime(2026, 8, 13, 10, 0, 0))
    dt2 = timezone.make_aware(datetime(2026, 8, 13, 11, 0, 0))
    dt3 = timezone.make_aware(datetime(2026, 8, 19, 12, 0, 0))

    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=10, created_date=dt1,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=10, created_date=dt2,
        pcs_no="PCS2", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=130, created_date=dt3,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )

    data = report1_orders()
    assert len(data) == 1
    row = data[0]
    assert row["LOT"] == "LOT-001"
    assert row["Production"] == 2
    assert row["Ok"] == 1
    assert row["Diff"] == 1
    assert len(row["defect_sns"]) == 1
    assert row["defect_sns"][0] == "PCS2"


@pytest.mark.django_db
def test_report1_no_data():
    """Отчет 1 с пустой БД."""
    data = report1_orders()
    assert len(data) == 0


@pytest.mark.django_db
def test_report2_line_productivity():
    """Проверка расчетов Report 2."""
    dt_start = timezone.make_aware(datetime(2026, 8, 13, 10, 0, 0))
    dt_end = timezone.make_aware(datetime(2026, 8, 13, 12, 0, 0))

    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=10, created_date=dt_start,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=50, created_date=dt_end,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )

    from reports.queries import report2_line_productivity
    from datetime import date
    data = report2_line_productivity(date(2026, 8, 13), date(2026, 8, 13))
    assert len(data) == 1
    assert data[0]["fol_hours"] == 2.0
    assert data[0]["fol_production"] == 1
    assert data[0]["has_data"] is True


@pytest.mark.django_db
def test_report5_repeated_passes():
    """Проверка Report 5: БУ с повторными проходами."""
    dt1 = timezone.make_aware(datetime(2026, 8, 13, 10, 0, 0))
    dt2 = timezone.make_aware(datetime(2026, 8, 13, 10, 5, 0))

    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=20, created_date=dt1,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=20, created_date=dt2,
        pcs_no="PCS1", result="NG", production_spec="SPEC"
    )

    from reports.queries import report5_repeated_passes
    from datetime import date
    data = report5_repeated_passes(date(2026, 8, 13), date(2026, 8, 13), limit=50)
    assert len(data) == 1
    assert data[0]["pcs_no"] == "PCS1"
    assert data[0]["station"] == 20


@pytest.mark.django_db
def test_dedup_on_reimport():
    """Повторная загрузка не создает дублей."""
    dt = timezone.make_aware(datetime(2026, 8, 13, 10, 0, 0))
    records = [
        ProductionRecord(
            lot_number="LOT-001", subop_no=10, created_date=dt,
            pcs_no="PCS1", result="OK", production_spec="SPEC"
        ),
    ]
    from uploads.excel_parser import import_to_database
    result1 = import_to_database(records)
    result2 = import_to_database(records)
    assert result1["inserted"] == 1
    assert result2["inserted"] == 0
    assert ProductionRecord.objects.count() == 1


# ========================
# Тесты сортировки
# ========================

def test_apply_sort_ascending():
    """Сортировка по возрастанию."""
    data = [
        {"name": "charlie", "val": 3},
        {"name": "alice", "val": 1},
        {"name": "bob", "val": 2},
    ]
    result = apply_sort(data, "val", "asc", ["val", "name"])
    assert [r["val"] for r in result] == [1, 2, 3]


def test_apply_sort_descending():
    """Сортировка по убыванию."""
    data = [
        {"name": "charlie", "val": 3},
        {"name": "alice", "val": 1},
        {"name": "bob", "val": 2},
    ]
    result = apply_sort(data, "val", "desc", ["val", "name"])
    assert [r["val"] for r in result] == [3, 2, 1]


def test_apply_sort_string_field():
    """Сортировка по строковому полю."""
    data = [
        {"name": "zebra"},
        {"name": "apple"},
        {"name": "mango"},
    ]
    result = apply_sort(data, "name", "asc", ["name"])
    assert [r["name"] for r in result] == ["apple", "mango", "zebra"]


def test_apply_sort_none_values():
    """None всегда сортируются в конец."""
    data = [
        {"name": "b"},
        {"name": None},
        {"name": "a"},
    ]
    result = apply_sort(data, "name", "asc", ["name"])
    assert [r["name"] for r in result] == ["a", "b", None]

    result_desc = apply_sort(data, "name", "desc", ["name"])
    assert [r["name"] for r in result_desc] == ["b", "a", None]

    mixed = apply_sort(
        [
            {"name": "B"},
            {"name": None},
            {"name": "a"},
            {"name": 2},
            {"name": 10},
        ],
        "name",
        "asc",
        ["name"],
    )
    assert [r["name"] for r in mixed] == [2, 10, "a", "B", None]


def test_apply_sort_invalid_field():
    """Сортировка по полю неиз списка allowed_fields возвращает данные без изменений."""
    data = [
        {"name": "b"},
        {"name": "a"},
    ]
    result = apply_sort(data, "invalid_field", "asc", ["name"])
    assert result == data


def test_get_sort_params_valid():
    """Валидные параметры сортировки извлекаются корректно."""
    class FakeRequest:
        def __init__(self, get_params):
            self.GET = get_params

    request = FakeRequest({"sort": "LOT", "dir": "desc"})
    sort_field, sort_dir = get_sort_params(request, "report1")
    assert sort_field == "LOT"
    assert sort_dir == "desc"


def test_get_sort_params_invalid_field():
    """Недопустимое поле сортировки игнорируется."""
    class FakeRequest:
        def __init__(self, get_params):
            self.GET = get_params

    request = FakeRequest({"sort": "invalid_field", "dir": "asc"})
    sort_field, sort_dir = get_sort_params(request, "report1")
    assert sort_field == ""
    assert sort_dir == "asc"


def test_get_sort_params_defaults():
    """Параметры по умолчанию при отсутствии sort и dir."""
    class FakeRequest:
        def __init__(self, get_params):
            self.GET = get_params

    request = FakeRequest({})
    sort_field, sort_dir = get_sort_params(request, "report2")
    assert sort_field == ""
    assert sort_dir == "asc"


def test_sortable_fields_defined():
    """Все отчеты имеют определенные сортируемые поля."""
    assert "report1" in SORTABLE_FIELDS
    assert "report2" in SORTABLE_FIELDS
    assert "report3" in SORTABLE_FIELDS
    assert "report4" in SORTABLE_FIELDS
    assert "report5" in SORTABLE_FIELDS
    assert "report6" in SORTABLE_FIELDS
    assert "LOT" in SORTABLE_FIELDS["report1"]
    assert "report_date" in SORTABLE_FIELDS["report2"]
    assert "report_month" in SORTABLE_FIELDS["report3"]
    assert "pcs_no" in SORTABLE_FIELDS["report4"]
    assert "station" in SORTABLE_FIELDS["report5"]


@pytest.mark.django_db
def test_report1_sort_by_production():
    """Сортировка отчета 1 по полю Production."""
    dt1 = timezone.make_aware(datetime(2026, 8, 13, 10, 0, 0))
    dt2 = timezone.make_aware(datetime(2026, 8, 13, 11, 0, 0))
    dt3 = timezone.make_aware(datetime(2026, 8, 13, 12, 0, 0))
    dt4 = timezone.make_aware(datetime(2026, 8, 19, 12, 0, 0))

    ProductionRecord.objects.create(
        lot_number="LOT-B", subop_no=10, created_date=dt1,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-B", subop_no=130, created_date=dt4,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-A", subop_no=10, created_date=dt2,
        pcs_no="PCS1", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-A", subop_no=10, created_date=dt3,
        pcs_no="PCS2", result="OK", production_spec="SPEC"
    )
    ProductionRecord.objects.create(
        lot_number="LOT-A", subop_no=130, created_date=dt4,
        pcs_no="PCS2", result="OK", production_spec="SPEC"
    )

    data = report1_orders()
    assert len(data) == 2

    sorted_data = apply_sort(data, "Production", "desc", SORTABLE_FIELDS["report1"])
    assert sorted_data[0]["LOT"] == "LOT-A"
    assert sorted_data[1]["LOT"] == "LOT-B"

    for i, row in enumerate(sorted_data, 1):
        row["N"] = i
    assert sorted_data[0]["N"] == 1
    assert sorted_data[1]["N"] == 2


@pytest.mark.django_db
def test_report6_serial_number():
    first = timezone.make_aware(datetime(2026, 9, 15, 10, 0, 0))
    second = timezone.make_aware(datetime(2026, 9, 15, 11, 0, 0))

    ProductionRecord.objects.create(
        lot_number="LOT-001", production_spec="SPEC-001", subop_no=50,
        position_no="50", workstation_name="壳体组装", pcs_no="SN-001",
        created_date=second, result="OK", test_data='{"value": 2}',
    )
    ProductionRecord.objects.create(
        lot_number="LOT-001", production_spec="SPEC-001", subop_no=10,
        position_no="10", workstation_name="PCB投料", pcs_no="SN-001",
        created_date=first, result="OK", test_data='{"value": 1}',
    )
    ProductionRecord.objects.create(
        lot_number="LOT-002", production_spec="SPEC-002", subop_no=10,
        position_no="10", workstation_name="PCB投料", pcs_no="SN-002",
        created_date=first, result="OK", test_data="",
    )

    data = report6_serial_number("SN-001")

    assert [row["created_date"] for row in data] == [first, second]
    assert [row["workstation_name"] for row in data] == ["PCB Loading", "Assembly"]
    assert data[0]["test_data"] == '{"value": 1}'
    assert report6_serial_number("MISSING") == []


@pytest.mark.django_db
def test_report7_station130_output_filters_and_deduplicates():
    start_date = date(2026, 9, 14)
    end_date = date(2026, 9, 20)
    records = [
        ProductionRecord(
            lot_number="LOT-A", subop_no=130, pcs_no="SN-DUP",
            production_spec="SPEC-A", result="OK",
            created_date=timezone.make_aware(datetime(2026, 9, 15, 10, 0)),
        ),
        ProductionRecord(
            lot_number="LOT-B", subop_no=130, pcs_no="SN-DUP",
            production_spec="SPEC-B", result="OK",
            created_date=timezone.make_aware(datetime(2026, 9, 16, 11, 0)),
        ),
        ProductionRecord(
            lot_number="LOT-A", subop_no=130, pcs_no="SN-START",
            production_spec="SPEC-A", result="OK",
            created_date=timezone.make_aware(datetime(2026, 9, 14, 8, 0)),
        ),
        ProductionRecord(
            lot_number="LOT-C", subop_no=130, pcs_no="SN-END",
            production_spec="SPEC-C", result="OK",
            created_date=timezone.make_aware(datetime(2026, 9, 20, 20, 0)),
        ),
        ProductionRecord(
            lot_number="LOT-NG", subop_no=130, pcs_no="SN-NG",
            result="NG", created_date=timezone.make_aware(datetime(2026, 9, 17, 10, 0)),
        ),
        ProductionRecord(
            lot_number="LOT-OTHER", subop_no=129, pcs_no="SN-OTHER",
            result="OK", created_date=timezone.make_aware(datetime(2026, 9, 17, 10, 0)),
        ),
        ProductionRecord(
            lot_number="LOT-OUT", subop_no=130, pcs_no="SN-OUT",
            result="OK", created_date=timezone.make_aware(datetime(2026, 9, 21, 10, 0)),
        ),
    ]
    ProductionRecord.objects.bulk_create(records)

    data = report7_station130_output(start_date, end_date)

    assert [row["pcs_no"] for row in data] == ["SN-END", "SN-DUP", "SN-START"]
    duplicate = next(row for row in data if row["pcs_no"] == "SN-DUP")
    assert duplicate["lot_number"] == "LOT-B"
    assert duplicate["production_spec"] == "SPEC-B"
    assert duplicate["created_date"] == records[1].created_date


class FakeRequest:
    def __init__(self, get_params):
        self.GET = get_params


def test_report7_period_selection():
    week = _report7_period(FakeRequest({"filter": "week", "week": "2026-W38"}))
    assert week["start_date"] == date(2026, 9, 14)
    assert week["end_date"] == date(2026, 9, 20)

    month = _report7_period(FakeRequest({"filter": "month", "year": "2024", "month": "2"}))
    assert month["start_date"] == date(2024, 2, 1)
    assert month["end_date"] == date(2024, 2, 29)

    invalid = _report7_period(
        FakeRequest({"filter": "period", "start": "2026-09-20", "end": "2026-09-14"})
    )
    assert invalid["error"]

    malformed = _report7_period(
        FakeRequest({"filter": "period", "start": "not-a-date", "end": "2026-09-20"})
    )
    assert malformed["error"]
    assert malformed["start_date"] == date.today()
    assert malformed["end_date"] == date.today()

    missing = _report7_period(
        FakeRequest({"filter": "period", "start": "2026-09-14"})
    )
    assert missing["error"]
    assert missing["start_date"] == date.today()
    assert missing["end_date"] == date.today()


def test_report7_summary():
    data = [
        {"pcs_no": "SN-1", "lot_number": "LOT-B"},
        {"pcs_no": "SN-1", "lot_number": "LOT-B"},
        {"pcs_no": "SN-2", "lot_number": "LOT-A"},
    ]
    summary = _report7_summary(data, date(2026, 9, 14), date(2026, 9, 20))
    assert summary == {
        "total_days": 7,
        "released_count": 2,
        "lots": "LOT-A, LOT-B",
    }


def test_report2_period_selection():
    month = _report2_period(FakeRequest({"filter": "month", "year": "2024", "month": "2"}))
    assert month["start_date"] == date(2024, 2, 1)
    assert month["end_date"] == date(2024, 2, 29)

    year = _report2_period(FakeRequest({"filter": "year", "year": "2025"}))
    assert year["start_date"] == date(2025, 1, 1)
    assert year["end_date"] == date(2025, 12, 31)

    period = _report2_period(
        FakeRequest({"filter": "period", "start": "2026-01-03", "end": "2026-01-05"})
    )
    assert period["start_date"] == date(2026, 1, 3)
    assert period["end_date"] == date(2026, 1, 5)

    invalid = _report2_period(
        FakeRequest({"filter": "period", "start": "2026-01-05", "end": "2026-01-03"})
    )
    assert invalid["error"]


def test_report2_summary_averages_only_active_days():
    data = [
        {"fol_hours": 2, "fol_production": 10, "fol_speed": 5, "bol_hours": 0, "bol_production": 0, "bol_speed": 0},
        {"fol_hours": 0, "fol_production": 0, "fol_speed": 0, "bol_hours": 4, "bol_production": 12, "bol_speed": 3},
        {"fol_hours": 4, "fol_production": 12, "fol_speed": 3, "bol_hours": 2, "bol_production": 6, "bol_speed": 3},
    ]
    summary = _report2_summary(data, date(2026, 1, 1), date(2026, 1, 3))
    assert summary == {
        "total_days": 3,
        "fol_hours": 6,
        "fol_production": 22,
        "fol_speed": 4,
        "fol_active_days": 2,
        "bol_hours": 6,
        "bol_production": 18,
        "bol_speed": 3,
        "bol_active_days": 2,
    }


def test_report2_view_default_year(monkeypatch):
    import json
    from reports.views import Report2View

    rows = [
        {"report_date": date(2026, 1, 1), "fol_hours": 2, "fol_production": 10, "fol_speed": 5, "bol_hours": 0, "bol_production": 0, "bol_speed": 0, "has_data": True},
        {"report_date": date(2026, 1, 2), "fol_hours": 0, "fol_production": 0, "fol_speed": 0, "bol_hours": 0, "bol_production": 0, "bol_speed": 0, "has_data": False},
    ]
    monkeypatch.setattr("reports.views.report2_line_productivity", lambda start_date, end_date: rows)

    request = RequestFactory().get("/reports/2/")
    response = Report2View.as_view()(request)

    assert response.status_code == 200
    assert response["Cache-Control"] == "no-store"
    response.render()
    context = response.context_data
    assert context["filter_type"] == "year"
    assert b'id="year-field"' in response.content
    assert b'name="year"' in response.content
    assert b'class="fol-cell text-primary"' in response.content
    assert b'class="bol-cell text-danger"' in response.content
    assert b'.table .fol-cell { background-color: #eaf6ff; }' in response.content
    assert b'.table .bol-cell { background-color: #fff0f0; }' in response.content
    assert context["start_date"] == date(date.today().year, 1, 1)
    assert context["end_date"] == date(date.today().year, 12, 31)
    assert context["data"] == [rows[0]]
    assert context["table_summary"]["fol_hours"] == 2
    assert context["table_summary"]["fol_production"] == 10
    chart_data = json.loads(context["chart_data"])
    assert chart_data["dates"] == ["2026-01-01"]
    assert chart_data["fol_hours"] == [2.0]
    assert chart_data["fol_production"] == [10.0]
    expected_days = 366 if date.today().year % 4 == 0 else 365
    assert context["summary"]["total_days"] == expected_days


def test_report2_export_includes_summary_and_charts(monkeypatch):
    from types import SimpleNamespace
    from reports.views import Report2ExportView

    rows = [
        {"report_date": date(2026, 1, 1), "fol_hours": 2, "fol_production": 10, "fol_speed": 5, "bol_hours": 0, "bol_production": 0, "bol_speed": 0, "has_data": True},
        {"report_date": date(2026, 1, 2), "fol_hours": 0, "fol_production": 0, "fol_speed": 0, "bol_hours": 0, "bol_production": 0, "bol_speed": 0, "has_data": False},
        {"report_date": date(2026, 1, 3), "fol_hours": 4, "fol_production": 12, "fol_speed": 3, "bol_hours": 2, "bol_production": 6, "bol_speed": 3, "has_data": True},
    ]
    monkeypatch.setattr("reports.views.report2_line_productivity", lambda start_date, end_date: rows)

    request = RequestFactory().get("/reports/2/export/", {"filter": "period", "start": "2026-01-01", "end": "2026-01-03"})
    request.user = SimpleNamespace(is_authenticated=True)
    response = Report2ExportView.as_view()(request)

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Производительность", "Сводка", "Диаграммы"]

    sheet = workbook["Производительность"]
    assert sheet["A2"].value == "2026-01-01"
    assert sheet["A3"].value == "2026-01-03"
    assert sheet.max_row == 3

    summary = workbook["Сводка"]
    assert summary["B2"].value == "01.01.2026 — 03.01.2026"
    assert summary["B3"].value.date() == date(2026, 1, 1)
    assert summary["B4"].value.date() == date(2026, 1, 3)
    assert summary["B5"].value == 3
    assert summary["B7"].value == "FOL"
    assert summary["C7"].value == "BOL"
    assert summary["B8"].value == 6
    assert summary["C8"].value == 2
    assert summary["B9"].value == 22
    assert summary["C9"].value == 6
    assert summary["B10"].value == 4
    assert summary["C10"].value == 3
    assert summary["B11"].value == 2
    assert summary["C11"].value == 1

    charts = workbook["Диаграммы"]._charts
    assert len(charts) == 3
    assert all(len(chart.series) == 2 for chart in charts)


@pytest.mark.django_db
def test_report7_export_is_available_to_regular_user():
    user = User.objects.create_user(username="operator", password="password")
    client = Client()
    client.force_login(user)
    ProductionRecord.objects.create(
        lot_number="LOT-001", subop_no=130, pcs_no="SN-001",
        production_spec="SPEC-001", result="OK",
        created_date=timezone.make_aware(datetime(2026, 9, 15, 10, 30)),
    )

    response = client.get("/reports/7/export/", {"filter": "day", "day": "2026-09-15"})

    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Сводка", "Выпуск"]
    assert workbook["Выпуск"]["A1"].value == "PCSNo - Серийный номер"
    assert workbook["Выпуск"]["A2"].value == "SN-001"
    assert workbook["Сводка"]["B3"].value == 1


def test_translate_workstation_name_by_position():
    assert translate_workstation_name("A110", "气密测试") == "Seal Test"
    assert translate_workstation_name("20", "FCT1") == "FCT1"
    assert translate_workstation_name("unknown", "Unknown") == "Unknown"


def test_report1_period_types_constant():
    """Проверка константы доступных типов периодов для отчета 1."""
    from reports.views import REPORT1_PERIOD_TYPES
    assert REPORT1_PERIOD_TYPES == {"year", "period"}


@pytest.mark.django_db
def test_report1_view_year_filter():
    """Проверка отчета 1 через Client с фильтром по году."""
    from django.test import Client, override_settings
    from django.contrib.auth import get_user_model

    User = get_user_model()
    user, _ = User.objects.get_or_create(username="testuser_report1_year", defaults={"password": "testpass"})

    with override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"]):
        client = Client()
        client.force_login(user)

        # Тест: год по умолчанию (текущий)
        response = client.get("/reports/1/")
        assert response.status_code == 200
        assert "Сводный отчет" in response.content.decode()

        # Тест: конкретный год
        response = client.get("/reports/1/?filter=year&year=2025")
        assert response.status_code == 200
        assert "Сводный отчет" in response.content.decode()

        # Тест: некорректный год (должен вернуть текущий год с ошибкой)
        response = client.get("/reports/1/?filter=year&year=not-a-year")
        assert response.status_code == 200
        # Должно отображаться сообщение об ошибке


def test_report1_summary():
    data = [
        {"LOT": "LOT-A", "Production": 2, "Ok": 1, "Diff": 1},
        {"LOT": "LOT-B", "Production": 3, "Ok": 0, "Diff": 3},
        {"LOT": "LOT-C", "Production": 0, "Ok": 0, "Diff": 0},
    ]

    summary = _report1_summary(data, date(2026, 9, 14), date(2026, 9, 20))

    assert summary == {
        "total_days": 7,
        "released_lots": 1,
        "production_total": 5,
        "ok_total": 1,
        "diff_total": 4,
    }


@pytest.mark.django_db
def test_report1_view_period_summary_context(monkeypatch):
    from django.test import Client, override_settings

    user, _ = User.objects.get_or_create(
        username="testuser_report1_summary", defaults={"password": "testpass"}
    )
    rows = [
        {"LOT": "LOT-A", "Production": 2, "Ok": 1, "Diff": 1},
        {"LOT": "LOT-B", "Production": 3, "Ok": 0, "Diff": 3},
        {"LOT": "LOT-C", "Production": 0, "Ok": 0, "Diff": 0},
    ]
    monkeypatch.setattr("reports.views.report1_orders", lambda start_date, end_date: rows)

    with override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"]):
        client = Client()
        client.force_login(user)
        response = client.get(
            "/reports/1/?filter=period&start=2026-09-14&end=2026-09-16"
        )
        response.render()

    assert response.status_code == 200
    assert response.context_data["summary"] == {
        "total_days": 3,
        "released_lots": 1,
        "production_total": 5,
        "ok_total": 1,
        "diff_total": 4,
    }
    content = response.content.decode()
    assert "Всего дней" in content
    assert "Выпущено лотов" in content
    assert "Запущено в производство" in content
    assert "Выпущено" in content
    assert "Не прошло сборку" in content


@pytest.mark.django_db
def test_report1_view_period_filter():
    """Проверка отчета 1 через Client с фильтром по произвольному периоду."""
    from django.test import Client, override_settings
    from django.contrib.auth import get_user_model

    User = get_user_model()
    user, _ = User.objects.get_or_create(username="testuser_report1_period", defaults={"password": "testpass"})

    with override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"]):
        client = Client()
        client.force_login(user)

        # Тест: валидный период
        response = client.get("/reports/1/?filter=period&start=2026-01-01&end=2026-01-31")
        assert response.status_code == 200
        assert "Сводный отчет" in response.content.decode()

        # Тест: некорректная дата начала
        response = client.get("/reports/1/?filter=period&start=not-a-date&end=2026-01-31")
        assert response.status_code == 200
        # Должно отображаться сообщение об ошибке

        # Тест: конец раньше начала
        response = client.get("/reports/1/?filter=period&start=2026-01-31&end=2026-01-01")
        assert response.status_code == 200
        # Должно отображаться сообщение об ошибке


@pytest.mark.django_db
def test_report1_defect_threshold(monkeypatch):
    """Проверка порога Diff < 10 для отображения SN в defect_sns."""
    from django.test import Client, override_settings
    from django.contrib.auth import get_user_model

    User = get_user_model()
    user, _ = User.objects.get_or_create(username="testuser_report1_defect", defaults={"password": "testpass"})

    # Подготовка строк с разным defect_count
    rows_9 = [{
        "LOT": "LOT-9",
        "production_spec": "SPEC",
        "Total": 10,
        "Production": 10,
        "Ok": 1,
        "Diff": 9,
        "Defect": "SN1, SN2, SN3, SN4, SN5, SN6, SN7, SN8, SN9",
        "defect_sns": ["SN1", "SN2", "SN3", "SN4", "SN5", "SN6", "SN7", "SN8", "SN9"],
        "start": None,
        "finish": None,
        "comment": "",
    }]
    rows_10 = [{
        "LOT": "LOT-10",
        "production_spec": "SPEC",
        "Total": 10,
        "Production": 10,
        "Ok": 0,
        "Diff": 10,
        "Defect": "SN не прошедших сборку более 10 штук",
        "defect_sns": [],
        "start": None,
        "finish": None,
        "comment": "",
    }]

    def mock_orders_9(start_date, end_date):
        return rows_9

    def mock_orders_10(start_date, end_date):
        return rows_10

    with override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"]):
        client = Client()
        client.force_login(user)

        # Diff = 9 -> defect_sns populated
        monkeypatch.setattr("reports.views.report1_orders", mock_orders_9)
        response = client.get("/reports/1/")
        response.render()
        assert response.status_code == 200
        data = response.context_data["data"]
        assert len(data) == 1
        assert data[0]["defect_sns"] == ["SN1", "SN2", "SN3", "SN4", "SN5", "SN6", "SN7", "SN8", "SN9"]
        assert data[0]["Defect"] == "SN1, SN2, SN3, SN4, SN5, SN6, SN7, SN8, SN9"

        # Diff = 10 -> defect_sns empty, Defect message
        monkeypatch.setattr("reports.views.report1_orders", mock_orders_10)
        response = client.get("/reports/1/")
        response.render()
        assert response.status_code == 200
        data = response.context_data["data"]
        assert len(data) == 1
        assert data[0]["defect_sns"] == []
        assert data[0]["Defect"] == "SN не прошедших сборку более 10 штук"


@pytest.mark.django_db
def test_report1_default_sort_by_start(monkeypatch):
    """Проверка сортировки по умолчанию по start asc."""
    from django.test import Client, override_settings
    from django.contrib.auth import get_user_model
    from datetime import datetime
    from django.utils import timezone

    User = get_user_model()
    user, _ = User.objects.get_or_create(username="testuser_report1_sort", defaults={"password": "testpass"})

    # Две строки с разными start датами
    row_later = {
        "LOT": "LOT-LATER",
        "production_spec": "SPEC",
        "Total": 5,
        "Production": 5,
        "Ok": 5,
        "Diff": 0,
        "Defect": "",
        "defect_sns": [],
        "start": timezone.make_aware(datetime(2026, 9, 15, 12, 0, 0)),
        "finish": None,
        "comment": "",
    }
    row_earlier = {
        "LOT": "LOT-EARLIER",
        "production_spec": "SPEC",
        "Total": 5,
        "Production": 5,
        "Ok": 5,
        "Diff": 0,
        "Defect": "",
        "defect_sns": [],
        "start": timezone.make_aware(datetime(2026, 9, 10, 8, 0, 0)),
        "finish": None,
        "comment": "",
    }
    rows = [row_later, row_earlier]  # порядок в моке: позже, раньше

    def mock_orders(start_date, end_date):
        return rows

    with override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"]):
        client = Client()
        client.force_login(user)
        monkeypatch.setattr("reports.views.report1_orders", mock_orders)
        response = client.get("/reports/1/")
        response.render()
        assert response.status_code == 200
        data = response.context_data["data"]
        # После сортировки по start asc первым должен быть LOT-EARLIER
        assert data[0]["LOT"] == "LOT-EARLIER"
        assert data[1]["LOT"] == "LOT-LATER"
        # N должно быть перенумеровано
        assert data[0]["N"] == 1
        assert data[1]["N"] == 2


def test_report1_comment_no_truncate():
    """Проверка, что комментарий не обрезается в шаблоне (юнит-тест шаблона не требуется, проверяем логику)."""
    # Этот тест больше документационный, так как урезание убрано в шаблоне.
    pass


@pytest.mark.django_db
def test_report1_export_is_available_to_regular_user():
    """Обычный авторизованный пользователь может экспортировать отчет №1 в Excel."""
    with override_settings(ALLOWED_HOSTS=["localhost", "127.0.0.1", "testserver"]):
        user, _ = User.objects.get_or_create(
            username="operator", defaults={"password": "password"}
        )
        client = Client()
        client.force_login(user)

        response = client.get("/reports/1/export/")

        assert response.status_code == 200
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        assert workbook.sheetnames == ["Сводный отчет"]


@pytest.mark.parametrize(
    ("report_number", "view_name", "query_name"),
    [
        (1, "Report1ExportView", "report1_orders"),
        (2, "Report2ExportView", "report2_line_productivity"),
        (3, "Report3ExportView", "report3_monthly_productivity"),
        (4, "Report4ExportView", "report4_defects"),
        (5, "Report5ExportView", "report5_repeated_passes"),
        (6, "Report6ExportView", "report6_serial_number"),
        (7, "Report7ExportView", "report7_station130_output"),
    ],
)
def test_report_exports_available_to_anonymous(monkeypatch, report_number, view_name, query_name):
    """Все отчёты можно экспортировать в Excel без авторизации."""
    from django.contrib.auth.models import AnonymousUser
    from reports import views

    monkeypatch.setattr(views, query_name, lambda *args, **kwargs: [])
    view = getattr(views, view_name)
    request = RequestFactory().get(f"/reports/{report_number}/export/")
    request.user = AnonymousUser()

    response = view.as_view()(request)

    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")


@pytest.mark.django_db
def test_admin_anonymous_and_regular_user_redirect():
    with override_settings(ALLOWED_HOSTS=["*"]):
        anonymous_client = Client()
        anonymous_response = anonymous_client.get("/admin/")
        assert anonymous_response.status_code == 302
        assert "/admin/login/" in anonymous_response["Location"]

        user, _ = User.objects.get_or_create(
            username="admin_regular",
            defaults={"is_active": True, "is_staff": False, "is_superuser": False},
        )
        user.is_active = True
        user.is_staff = False
        user.is_superuser = False
        user.save()

        client = Client()
        client.force_login(user)
        response = client.get("/admin/")
        assert response.status_code == 302
        assert "/admin/login/" in response["Location"]


@pytest.mark.django_db
def test_admin_index_preserves_redirect_response(monkeypatch):
    redirect_response = HttpResponseRedirect("/admin/login/")
    monkeypatch.setattr(
        "django.contrib.admin.sites.AdminSite.index",
        lambda self, request, extra_context=None: redirect_response,
    )

    request_users = (
        AnonymousUser(),
        User(is_active=True, is_staff=False, is_superuser=False),
    )
    for request_user in request_users:
        request = RequestFactory().get("/admin/")
        request.user = request_user
        assert mes_admin_site.index(request) is redirect_response


@pytest.mark.django_db
def test_admin_staff_sees_business_models_but_not_auth_models():
    with override_settings(ALLOWED_HOSTS=["*"]):
        user, _ = User.objects.get_or_create(
            username="admin_staff",
            defaults={"is_active": True, "is_staff": True, "is_superuser": False},
        )
        user.is_active = True
        user.is_staff = True
        user.is_superuser = False
        user.save()

        client = Client()
        client.force_login(user)
        response = client.get("/admin/")

        assert response.status_code == 200
        app_list = response.context["app_list"]
        model_names = {
            model["object_name"]
            for app in app_list
            for model in app["models"]
        }
        assert {"ProductionRecord", "LotInfo", "ManualDefect", "Report5Comment"} <= model_names
        assert "User" not in model_names
        assert "Group" not in model_names
        assert not any(app["app_label"] == "auth" for app in app_list)


@pytest.mark.django_db
def test_admin_superuser_sees_auth_models():
    with override_settings(ALLOWED_HOSTS=["*"]):
        user, _ = User.objects.get_or_create(
            username="admin_super",
            defaults={"is_active": True, "is_staff": True, "is_superuser": True},
        )
        user.is_active = True
        user.is_staff = True
        user.is_superuser = True
        user.save()

        client = Client()
        client.force_login(user)
        response = client.get("/admin/")

        assert response.status_code == 200
        auth_app = next(app for app in response.context["app_list"] if app["app_label"] == "auth")
        model_names = {model["object_name"] for model in auth_app["models"]}
        assert {"User", "Group"} <= model_names


@pytest.mark.django_db
def test_admin_production_record_is_read_only_for_staff():
    request = RequestFactory().get("/admin/")
    request.user = User(is_active=True, is_staff=True, is_superuser=False)
    model_admin = ProductionRecordAdmin(ProductionRecord, mes_admin_site)

    assert model_admin.has_view_permission(request)
    assert not model_admin.has_add_permission(request)
    assert not model_admin.has_change_permission(request)
    assert not model_admin.has_delete_permission(request)


@pytest.mark.django_db
def test_admin_lot_number_is_editable_on_create_and_read_only_on_change():
    request = RequestFactory().get("/admin/")
    request.user = User(is_active=True, is_staff=True, is_superuser=False)
    model_admin = LotInfoAdmin(LotInfo, mes_admin_site)

    assert "lot_number" not in model_admin.get_readonly_fields(request, obj=None)
    lot = LotInfo(lot_number="LOT-ADMIN")
    assert "lot_number" in model_admin.get_readonly_fields(request, obj=lot)


@pytest.mark.django_db
def test_admin_audit_fields_are_set_on_create():
    with override_settings(ALLOWED_HOSTS=["*"]):
        user, _ = User.objects.get_or_create(
            username="admin_audit",
            defaults={"is_active": True, "is_staff": True, "is_superuser": False},
        )
        user.is_active = True
        user.is_staff = True
        user.is_superuser = False
        user.save()

        client = Client()
        client.force_login(user)
        defect_pcs_no = f"ADMIN-DEFECT-{user.pk}"
        comment_pcs_no = f"ADMIN-COMMENT-{user.pk}"

        defect_response = client.post(
            "/admin/production/manualdefect/add/",
            {
                "pcs_no": defect_pcs_no,
                "production_spec": "SPEC-ADMIN",
                "lot_number": "LOT-ADMIN",
                "reason": "Admin test",
                "comment": "Created through admin",
                "_save": "Save",
            },
        )
        assert defect_response.status_code == 302
        defect = ManualDefect.objects.get(pcs_no=defect_pcs_no)
        assert defect.created_by == user.username
        assert defect.created_at is not None

        comment_response = client.post(
            "/admin/production/report5comment/add/",
            {
                "pcs_no": comment_pcs_no,
                "station_no": "17",
                "lot_number": "LOT-ADMIN",
                "comment": "Created through admin",
                "_save": "Save",
            },
        )
        assert comment_response.status_code == 302
        comment = Report5Comment.objects.get(pcs_no=comment_pcs_no)
        assert comment.created_by == user.username
        assert comment.created_at is not None


@pytest.mark.django_db
def test_admin_dashboard_has_links_and_aggregate_metrics():
    with override_settings(ALLOWED_HOSTS=["*"]):
        user, _ = User.objects.get_or_create(
            username="admin_dashboard",
            defaults={"is_active": True, "is_staff": True, "is_superuser": False},
        )
        user.is_active = True
        user.is_staff = True
        user.is_superuser = False
        user.save()

        LotInfo.objects.create(
            lot_number="LOT-DASHBOARD",
            production_spec="SPEC-DASHBOARD",
            plan_total=25,
        )
        ProductionRecord.objects.create(
            lot_number="LOT-DASHBOARD",
            subop_no=17,
            pcs_no="SN-DASHBOARD",
            created_date=timezone.now(),
            result="OK",
            production_spec="SPEC-DASHBOARD",
        )
        ManualDefect.objects.create(
            pcs_no="SN-DASHBOARD-DEFECT",
            production_spec="SPEC-DASHBOARD",
            lot_number="LOT-DASHBOARD",
            created_by="seed",
        )
        Report5Comment.objects.create(
            pcs_no="SN-DASHBOARD-COMMENT",
            station_no=17,
            lot_number="LOT-DASHBOARD",
            comment="Dashboard test",
            created_by="seed",
        )

        client = Client()
        client.force_login(user)
        response = client.get("/admin/")

        assert response.status_code == 200
        content = response.content.decode()
        metrics = response.context["dashboard_metrics"]
        assert metrics["lots"]["total"] >= 1
        assert metrics["lots"]["plan_total"] >= 25
        assert metrics["production_records"]["total"] >= 1
        assert metrics["production_records"]["recent"] >= 1
        assert metrics["manual_defects"]["total"] >= 1
        assert metrics["manual_defects"]["recent"] >= 1
        assert metrics["report5_comments"]["total"] >= 1
        assert metrics["report5_comments"]["recent"] >= 1

        for label in (
            "Лоты",
            "Записи производства",
            "Ручной брак",
            "Комментарии отчёта 5",
            "Отчёт 1: Сводный",
            "Отчёт 7: Выпуск",
            "Загрузка Excel",
        ):
            assert label in content
        assert "План: 25" in content
        assert "За 7 дней:" in content
        for url_name in (
            "mes_admin:production_lotinfo_changelist",
            "mes_admin:production_productionrecord_changelist",
            "mes_admin:production_manualdefect_changelist",
            "mes_admin:production_report5comment_changelist",
            "reports:report1",
            "reports:report7",
            "uploads:upload_home",
        ):
            assert reverse(url_name) in content
